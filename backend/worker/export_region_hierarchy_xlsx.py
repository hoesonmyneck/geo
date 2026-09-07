"""Выгрузка населения региона в Excel тремя листами по уровням иерархии:
    лист 1  Область -> Район
    лист 2  Область -> Район -> Сельский округ
    лист 3  Область -> Район -> Сельский округ -> Населённый пункт

Источник — pop_dwelling (жилища с демографией в stats). Уровня «сельский округ»
в pop_dwelling нет: берём его из corpus.full_path_rus (адреса egov содержат
«СЕЛЬСКИЙ ОКРУГ X») через общие РКА. У городов этого уровня нет по определению,
им проставляем «— (город)», чтобы суммы всех трёх листов сходились.

Запуск (ХОСТ):
    python backend/worker/export_region_hierarchy_xlsx.py "КАРАГАНДИНСКАЯ" [out.xlsx]
"""
from __future__ import annotations
import sys
import psycopg
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

DSN = "host=localhost port=5432 dbname=geo user=geo password=geopassword123"

# Одна база для всех трёх листов: жилище + его место в иерархии.
BASE = r"""
WITH d AS (
    SELECT dwelling_id, kind, geocode_addr,
           initcap(stats->>'regname') AS obl,
           initcap(stats->>'rainame') AS rai,
           (stats->>'total')::int     AS total
      FROM pop_dwelling
     WHERE stats->>'regname' ILIKE %(reg)s
), ok AS (
    -- сельский округ: самый частый по адресам egov этого села
    SELECT m.dwelling_id,
           mode() WITHIN GROUP (
               ORDER BY initcap(btrim((regexp_match(c.full_path_rus,
                        'СЕЛЬСКИЙ ОКРУГ\s+([^,]+)', 'i'))[1]))
           ) AS okrug
      FROM d
      JOIN pop_rca_dwelling m ON m.dwelling_id = d.dwelling_id
      JOIN corpus c           ON c.rca = m.rca
     WHERE d.kind = 'village' AND c.full_path_rus ~* 'СЕЛЬСКИЙ ОКРУГ'
     GROUP BY m.dwelling_id
), h AS (
    SELECT d.obl, d.rai, d.total,
           CASE WHEN d.kind = 'village'
                -- у посёлков/разъездов уровня «сельский округ» нет и в egov
                -- (район -> посёлок напрямую) — это не пропуск данных
                THEN coalesce(o.okrug,
                       CASE WHEN d.geocode_addr ~* '^\s*(ПОСЕЛОК|ПОСЁЛОК|РАЗЪЕЗД|СТАНЦИЯ)'
                            THEN '— (посёлок, округа нет)'
                            ELSE '— (округ не определён)' END)
                ELSE '— (город)' END AS okrug,
           CASE WHEN d.kind = 'village'
                THEN coalesce(nullif(initcap(btrim(split_part(d.geocode_addr, ',', 1))), ''),
                              '— (адрес не указан)')
                ELSE coalesce('Город ' || initcap(btrim((regexp_match(d.geocode_addr,
                       '(?:ГОРОД\s+(?:ОБЛАСТНОГО|РАЙОННОГО|РЕСПУБЛИКАНСКОГО)\s+ЗНАЧЕНИЯ|ГОРОД)\s+([^,]+)',
                       'i'))[1])), '— (город не распознан)') END AS np
      FROM d LEFT JOIN ok o ON o.dwelling_id = d.dwelling_id
)
"""

SHEETS = [
    ("Районы", ["Область", "Район", "Количество людей"],
     BASE + "SELECT obl, rai, sum(total) FROM h GROUP BY 1,2 ORDER BY 1,2"),
    ("Сельские округа", ["Область", "Район", "Сельский округ", "Количество людей"],
     BASE + "SELECT obl, rai, okrug, sum(total) FROM h GROUP BY 1,2,3 ORDER BY 1,2,3"),
    ("Населённые пункты", ["Область", "Район", "Сельский округ", "Населённый пункт", "Количество людей"],
     BASE + "SELECT obl, rai, okrug, np, sum(total) FROM h GROUP BY 1,2,3,4 ORDER BY 1,2,3,4"),
]

WIDTHS = {"Область": 26, "Район": 28, "Сельский округ": 30,
          "Населённый пункт": 34, "Количество людей": 18}


def main() -> None:
    if len(sys.argv) < 2:
        print('Укажи регион, напр.: python ... "КАРАГАНДИНСКАЯ"', file=sys.stderr)
        sys.exit(1)
    region = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else f"население_{region.lower()}.xlsx"

    conn = psycopg.connect(DSN)
    wb = Workbook()
    wb.remove(wb.active)
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="2563EB")

    for title, header, sql in SHEETS:
        rows = conn.execute(sql, {"reg": f"%{region}%"}).fetchall()
        ws = wb.create_sheet(title)
        ws.append(header)
        for i, name in enumerate(header, start=1):
            ws.column_dimensions[get_column_letter(i)].width = WIDTHS[name]
        for cell in ws[1]:
            cell.font, cell.fill = hdr_font, hdr_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in rows:
            ws.append(list(r))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"

        # строка «Всего»
        last_col = get_column_letter(len(header))
        tr = len(rows) + 2
        ws.cell(tr, len(header) - 1, "ВСЕГО").font = Font(bold=True)
        c = ws.cell(tr, len(header), f"=SUM({last_col}2:{last_col}{tr - 1})")
        c.font = Font(bold=True)

        total = sum(r[-1] or 0 for r in rows)
        print(f"  лист «{title}»: строк {len(rows):,}, людей {total:,}", flush=True)

    conn.close()
    wb.save(out)
    print(f"\nГотово: {out}")


if __name__ == "__main__":
    main()
